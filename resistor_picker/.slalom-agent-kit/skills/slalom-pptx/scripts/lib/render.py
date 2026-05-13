"""Render the references/ tree from SLIDE_MAP + extracted inventories.

Produces:
    references/index.md             (small router)
    references/index.json           (machine-readable catalog)
    references/colors.md            (brand colors + fonts + footer)
    references/workflows.md         (A/B/C workflows + icon-swap + validation)
    references/icons.md             (icon library catalog)
    references/slides/<cat>.md      (one per category in CATEGORIES)

`colors.md` and `workflows.md` are static prose; `icons.md` and the per-category
slide catalogs are derived from inventories + slide_map.py + the icons folder.

Brand color values are read from ppt/theme/theme1.xml inside the master template so
they stay accurate if Slalom updates the palette.
"""

from __future__ import annotations

import hashlib
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Iterable

from slide_map import (
    CATEGORIES,
    CATEGORY_LABELS,
    SLIDE_MAP,
    SlideEntry,
    by_category,
)


COLOR_NAMES: dict[str, tuple[str, str]] = {
    "dk1": ("Black", "Text/foreground"),
    "lt1": ("White", "Background"),
    "dk2": ("Dark Gray", "Secondary text"),
    "lt2": ("Light Gray", "Secondary background"),
    "accent1": ("Slalom Blue", "Primary accent"),
    "accent2": ("Dark Blue", "Secondary accent"),
    "accent3": ("Cyan", "Accent 3"),
    "accent4": ("Coral Red", "Accent 4"),
    "accent5": ("Purple", "Accent 5"),
    "accent6": ("Chartreuse", "Accent 6"),
    "hlink": ("Slalom Blue", "Hyperlink"),
    "folHlink": ("Slalom Blue", "Followed hyperlink"),
}


def compute_master_sha256(master_pptx: Path) -> str:
    h = hashlib.sha256()
    with master_pptx.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def read_brand_colors(master_pptx: Path) -> list[tuple[str, str, str, str]]:
    """Return [(scheme, hex, slalom_name, role), ...] from the master theme."""
    with zipfile.ZipFile(master_pptx) as zf:
        with zf.open("ppt/theme/theme1.xml") as f:
            xml = f.read().decode()
    block = re.search(r"<a:clrScheme[^>]*>(.*?)</a:clrScheme>", xml, re.DOTALL)
    if not block:
        return []
    out = []
    for scheme, val in re.findall(
        r"<a:(\w+)>\s*<a:s[yr][sg]bClr[^>]*val=\"([^\"]+)\"", block.group(1)
    ):
        slalom_name, role = COLOR_NAMES.get(scheme, (scheme, ""))
        out.append((scheme, val.upper(), slalom_name, role))
    return out


def shape_label(shape: dict) -> str:
    """Compact one-token label for an inventory shape: placeholder type or first text."""
    ph = shape.get("placeholder_type")
    paras = shape.get("paragraphs", [])
    text = (paras[0].get("text") if paras else "") or ""
    if ph:
        return ph.lower()
    text = text.strip()
    if not text:
        return "shape"
    truncated = text[:30] + ("..." if len(text) > 30 else "")
    return f'"{truncated}"'


_PLACEHOLDER_NOISE_PREFIXES = ("click to add", "click here to add")
_PLACEHOLDER_NOISE_SUBSTRINGS = ("INSTRUCTIONS",)


def _is_placeholder_noise(shape: dict) -> bool:
    """Identify master-template help-text that's not actually editable content.

    Slalom's master uses 'INSTRUCTIONS' suffixes and 'Click to add ...' boilerplate
    on some slides as designer guidance; the inventory captures these as text
    shapes but agents shouldn't treat them as fillable slots. Filtering at the
    summary layer keeps the per-category catalogs honest.
    """
    if shape.get("placeholder_type"):
        return False  # Real placeholders (TITLE, BODY, etc.) are always kept.
    paras = shape.get("paragraphs", [])
    text = (paras[0].get("text") if paras else "") or ""
    text = text.strip()
    if not text:
        return False
    if any(s in text for s in _PLACEHOLDER_NOISE_SUBSTRINGS):
        return True
    lower = text.lower()
    return any(lower.startswith(p) for p in _PLACEHOLDER_NOISE_PREFIXES)


def summarise_shapes(inventory: dict) -> list[str]:
    """Return a per-type-counted shape summary in first-appearance order.

    Example: ['title', 'body', '4x "Subheadline"', '"Diagram title"', 'date', 'footer']
    even when "Subheadline" and "Diagram title" interleave in the underlying XML.

    Skips master-template designer-help text ('INSTRUCTIONS' / 'Click to add ...').
    """
    counts: dict[str, int] = {}
    order: list[str] = []
    for slide in inventory.values():
        for shape in slide.values():
            if _is_placeholder_noise(shape):
                continue
            label = shape_label(shape)
            if label not in counts:
                order.append(label)
            counts[label] = counts.get(label, 0) + 1

    return [f"{counts[l]}x {l}" if counts[l] > 1 else l for l in order]


def load_inventory(inv_path: Path) -> dict:
    if not inv_path.exists():
        return {}
    return json.loads(inv_path.read_text())


def render_colors_md(brand_colors: list[tuple[str, str, str, str]]) -> str:
    rows = "\n".join(
        f"| {scheme} | #{hex_val} | {name} | {role} |"
        for scheme, hex_val, name, role in brand_colors
    )
    return f"""# Brand Colors, Fonts, Footer

Use `<a:schemeClr val="accentN"/>` references in OOXML so colors track theme overrides.
Use Arial as a web-safe fallback when PptxGenJS can't load Avenir Next LT Pro.

## Color Scheme

| Scheme | Hex | Slalom Name | Role |
|---|---|---|---|
{rows}

## Fonts

- Body: theme font `+mn-lt` (Avenir Next LT Pro, Arial fallback)
- Headings: theme font `+mj-lt` (Avenir Next LT Pro, Arial fallback)

## Footer

Standard footer text on all body slides:

> Slalom. All Rights Reserved. Proprietary and Confidential.

## Date format

Default format for the system `DATE` placeholder on body slides: **`yyyy`** (4 chars; e.g. `2026`). The master's `DATE` placeholder is only ~0.28" wide — wider formats (`dd/mm/yyyy`, `mm/yy`, etc.) wrap onto two lines or overflow into the footer. The master ships `2025` as the stale default; replace with the actual deck-delivery year.

Cover slides use a different placeholder (a regular `BODY`) for any month/year eyebrow line — the `yyyy` rule applies only to the system `DATE` placeholder on body slides.

## DATE / FOOTER as a composite footer line

`DATE` (`x≈1.51"`) and `FOOTER` (`x=1.80"`) are two independent placeholders abutted at a fixed 0.01" gap. There's no shape grouping anchoring them, so they look like one continuous footer line *only* when the DATE content stays at ≤4 characters. Anything longer breaks the visual continuity. Treat them as a composite: keep DATE short (default `yyyy`) and let FOOTER carry the brand string.
"""


def render_workflows_md() -> str:
    return """# Workflows

| Need | Workflow |
|---|---|
| Layout already in catalog | A: Template Replace |
| Close match, need to add/remove shapes | C: Faithful OOXML |
| Combine elements from multiple slides | C: Faithful OOXML |
| New layout that must match Slalom style | C: Faithful OOXML |
| Internal / throwaway content; user explicitly waived Slalom branding fidelity | B: PptxGenJS |
| Confirm the output (run after every build) | D: Verify |

## Workflow A: Template Replace

```bash
# 1. Extract the slide(s) from the master by 0-based master idx
#    (look up idx in references/slides/<category>.md)
python3 "$PPTX_DIR/legacy/rearrange.py" \\
    "assets/master/Slalom Template document.pptx" working.pptx <idx>
# Multi-slide deck: pass a comma-separated list, e.g. 13,32,45,118,167

# 2. Read the inventory
cat assets/inventories/<slide-name>.json

# 3. Write replacement-text.json matching the inventory structure
#    (preserve font_size, bold, line_spacing from the original)
#    REQUIRED on every body slide - see "Required system-placeholder fills" below.

# 4. Apply replacements
python3 "$PPTX_DIR/legacy/replace.py" working.pptx replacement-text.json output.pptx

# 5. (Optional) Swap icons - see Icon Swap below.
```

`replacement-text.json` mirrors the inventory's `slide-N` / `shape-M` / `paragraphs[i].text` keys. Replace only what you want changed; **omitted entries keep the original master text — empty string `""` BLANKS the placeholder and is destructive.** When in doubt, omit.

### Required system-placeholder fills (body slides)

Every body / content / section-divider / closer slide carries three system placeholders that the master ships with stale or hardcoded defaults. Override `DATE` and `FOOTER` on every such slide; `SLIDE_NUMBER` is auto-filled by `replace.py` when omitted (see below). Cover slides typically don't carry these; check the inventory to confirm.

| Inventory `placeholder_type` | Master ships with | What to put |
|---|---|---|
| `DATE` | "2025" (stale) | Deck year, default format `yyyy` (e.g. "2026"). Master placeholder is only ~0.28" wide; longer formats wrap or clip. |
| `FOOTER` | brand string (often blanked accidentally) | `Slalom. All Rights Reserved. Proprietary and Confidential.` (from `references/colors.md`) |
| `SLIDE_NUMBER` | hardcoded master position (e.g. "29", "94") - NOT a dynamic field | Auto-filled with the slide's deck position (1-based) when omitted from your replacement JSON. Pass an explicit override only when you need a different number (e.g. `"iv"` for front-matter Roman numerals). |

After `replace.py`, run Workflow D1 (`markitdown`) and visually scan that the date, footer, and page numbers appear correctly on every body slide.

```json
{
  "slide-0": {
    "shape-2": {
      "paragraphs": [
        {"text": "Acme Q4 Strategy Review"}
      ]
    },
    "shape-3": {
      "paragraphs": [
        {"text": "Prepared for the executive team, January 2026."}
      ]
    }
  }
}
```

## On-demand thumbnails

Thumbnails are not pre-baked - render them when you actually need a visual.

```bash
# Pre-synthesis: preview a master layout before committing
TMP=$(mktemp -d) && trap 'rm -rf "$TMP"' EXIT
python3 "$PPTX_DIR/legacy/rearrange.py" \\
    "assets/master/Slalom Template document.pptx" "$TMP/preview.pptx" <idx>
python3 "$PPTX_DIR/modern/thumbnail.py" "$TMP/preview.pptx" "$TMP/thumb" --cols 1
# -> $TMP/thumb.jpg

# Post-synthesis: visual QA on the deck you just built (also Workflow D3)
python3 "$PPTX_DIR/modern/thumbnail.py" output.pptx /tmp/qa --cols 3
# -> /tmp/qa.jpg (grid)
```

Both paths use the vendored `modern/thumbnail.py` (driven by LibreOffice). Requires `soffice` on PATH; if missing, fall back to the markitdown text check (D1) and shape-type dump (D2).

## Workflow B: Hybrid Build (PptxGenJS)

1. Build slides programmatically with PptxGenJS.
2. Pull brand colors from `references/colors.md` (NO `#` prefix in PptxGenJS hex values).
3. Use Arial font (Avenir Next LT Pro is not web-safe).
4. Use pre-extracted icons via `addImage({ path: 'assets/icons/<name>.png' })`.
5. Add the standard footer text from `references/colors.md`.

## Workflow C: Faithful New Creation (OOXML)

Workflow C uses the modern `office/` toolchain: smart-quote auto-escape on unpack, schema-validating pack with `--original` comparison, and `clean.py` for orphan removal. Pack auto-validates — no separate validate step.

```bash
# 1. Extract closest template slide as base (legacy: deterministic by 0-based idx)
python3 "$PPTX_DIR/legacy/rearrange.py" \\
    "assets/master/Slalom Template document.pptx" base.pptx <idx>

# 2. Unpack (modern: auto-escapes smart quotes -> XML entities)
python3 "$PPTX_DIR/modern/unpack.py" base.pptx unpacked/

# 3. Edit ppt/slides/slide1.xml:
#    - Add/remove/reposition <p:sp> shape elements
#    - Use scheme color refs: <a:schemeClr val="accent1"/> (NOT hardcoded hex)
#    - Use theme fonts: <a:latin typeface="+mn-lt"/> (NOT explicit font names)
#    - Copy shapes from other unpacked template slides as building blocks
#    - Smart quotes are already escaped — write literal " ' in your edits and pack will round-trip them

# 4. (Optional) Drop orphans if you removed slides or relationships
python3 "$PPTX_DIR/modern/clean.py" unpacked/

# 5. Repack with validation against the original
python3 "$PPTX_DIR/modern/pack.py" unpacked/ output.pptx --original base.pptx
```

Starting from a template slide inherits master styles, theme colors, fonts, shadows, and gradients automatically. `pack.py --original` auto-repairs minor schema issues and reports whether the result validates.

## Icon Swap

```bash
# 1. Unpack the slide
python3 "$PPTX_DIR/modern/unpack.py" output.pptx unpacked/

# 2. Find which media files the slide's icons reference
cat unpacked/ppt/slides/_rels/slide1.xml.rels | grep media

# 3. Copy replacement icons (both PNG and SVG)
cp assets/icons/<new-icon>.png unpacked/ppt/media/<original-image>.png
cp assets/icons/<new-icon>.svg unpacked/ppt/media/<original-image>.svg

# 4. Repack
python3 "$PPTX_DIR/modern/pack.py" unpacked/ output.pptx --original output.pptx
```

## Workflow D: Verify

After every build, run all three checks. Each catches a different failure mode.

### D1. Text content (deterministic, fast)

```bash
python3 -m markitdown output.pptx
```

Confirms expected text appears. Fastest check.

### D2. Shape structure (deterministic, no rendering)

```bash
python3 - <<'PY'
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
prs = Presentation("output.pptx")
for i, slide in enumerate(prs.slides):
    print(f"slide {i+1}:")
    for sh in slide.shapes:
        tag = sh.shape_type.name if sh.shape_type else "?"
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            tag += f"({sh.auto_shape_type.name})"
        text = sh.text_frame.text[:40].replace("\\n", " / ") if sh.has_text_frame else ""
        print(f"  {tag:<28} {text!r}")
PY
```

Confirms the right shape *types* are present (e.g., a real `CHEVRON` auto-shape vs. just a rectangle). Inventory JSON only captures text shapes, so this is the structural ground truth.

### D3. Visual rendering (catches overflow, cutoff, alignment)

```bash
python3 "$PPTX_DIR/modern/thumbnail.py" output.pptx /tmp/qa --cols 3
```

Renders the deck to JPG via LibreOffice. Open the result to eyeball overflow, cutoff, font fallback, alignment.

### Other checks
- Slide count matches expectations.
- If `replace.py` errors on overflow, shorten the text and retry.
- After icon swap, verify replacement files are > 0 bytes.

## Verification gate

The deck is NOT ready until D1, D2, AND D3 have all been run AND inspected, and the agent has cited what each step found. Implicit success ("looks good", "tests passed") is not acceptable; the agent must report concrete evidence per check. Examples of acceptable evidence:

- **D1 (markitdown):** "title text 'Why we built X' is present on slide 2; no `_x000B_` or other escape leakage anywhere; FOOTER reads exactly 'Slalom. All Rights Reserved. Proprietary and Confidential.' on every body slide."
- **D2 (shape-type dump):** "slide 4 contains 4× CHEVRON auto-shapes (not rectangles); SLIDE_NUMBER placeholders show '1' through '6' across body slides 1-6."
- **D3 (thumbnail render):** "thumb.jpg shows no overflow in any title; chevron arrows point right; device-frame body copy stays clear of the iPhone overlay; closer slide imagery matches the deck context."

If `soffice` is unavailable and D3 cannot run, say so explicitly and flag the deck as **unverified** rather than skipping silently. D1 + D2 alone do not constitute a passed gate.

Iterate as needed. Common signals to revise: D3 reveals overflow that the deterministic overflow-checks missed (LibreOffice rendering differs from PIL estimation), the wrong placeholder type was filled (D2 caught it via shape-type dump), or text disagrees between D1 and D3 (escape leakage, encoding issue). The loop is designed: edit JSON → `replace.py` → D1+D2+D3 → adjust → repeat. Don't stop after one pass; trust the triad to reveal what each individual check misses.

## Pitfalls

- **Empty string is destructive.** Passing `{"text": ""}` for a shape in your replacement JSON BLANKS it. To inherit master text, OMIT the shape entirely; to override, supply real text. Omission ≠ empty string.
- **DATE / FOOTER on body slides ship with stale or hardcoded defaults from the master** (year "2025", brand footer string sometimes blanked). Always override them — see *Required system-placeholder fills* in Workflow A. Default DATE format is `yyyy` (4 chars; the master's DATE placeholder is only ~0.28" wide). SLIDE_NUMBER is auto-filled with the slide's deck position when omitted; pass an explicit override only for non-numeric formats.
- **Soft line breaks in master text use `\\x0b` (vertical tab).** Inventory JSON shows them as `\\u000b` between words where the master had a `<a:br/>`. `replace.py` (Patch #8) splits on `\\x0b` and emits real OOXML `<a:br/>` elements automatically — write the literal VT in your replacement text where you want a soft break. Older versions of the toolchain wrote the VT into `<a:t>` directly; PowerPoint serialized it as the literal escape `_x000B_` and rendered that string visibly. If you see `_x000B_` in a D1 markitdown output, the build is using a stale `replace.py` predating the patch.
- **Closer-slide stock imagery.** Layouts like `thank-you-1` and several `quote-N` variants ship with master-provided stock photos (winding road, generic forest, etc.). The inventory exposes only the text placeholders; the picture is a master-owned shape. To swap the image for context-appropriate photography, drop into Workflow C: unpack → replace `ppt/media/imageN.{jpg,png}` → pack. Same procedure as the Icon Swap above, just on a different media file.
- Delete `INSTRUCTIONS` and `Click to add ...` placeholder text in extracted slides; it's master designer-help, not editable content. The auto-derived `Key shapes` columns already hide it, but the actual `.pptx` files retain it.
- `bullets-4-cols.pptx` is a four-column-bullets layout, not a chevron, despite the master title saying "Chevron process". The real chevron is `chevron-4.pptx`.
- Inventory captures only text shapes. Icons are picture shapes; swap them via the Icon Swap procedure above.
- Cover slides use `CENTER_TITLE` / `BODY` placeholders (no `TITLE`); section dividers may have only `BODY`. Read the inventory's `placeholder_type` field; don't assume `TITLE` exists.

## Tips

- Always start from a template slide when visual fidelity matters.
- Prefer `<a:schemeClr val="accent1"/>` over `<a:srgbClr val="0C62FB"/>` in OOXML.
- Use `<a:latin typeface="+mn-lt"/>` for body, `+mj-lt` for headings.
- PptxGenJS builds outside the template theme - use only for drafts.
- Text must fit shape dimensions - check width/height in inventory JSON.
- Icons are PNG+SVG pairs - replace both when swapping.
- The slide master provides consistent footer, page numbers, and background.
- Combine slides via `rearrange.py` with multiple indices: `rearrange.py master.pptx out.pptx 13,104,92`.
"""


def render_icons_md(icons_dir: Path) -> str:
    pngs = sorted(p.stem for p in icons_dir.glob("*.png"))
    rows = []
    for stem in pngs:
        svg_present = (icons_dir / f"{stem}.svg").exists()
        svg_cell = f"icons/{stem}.svg" if svg_present else "-"
        rows.append(f"| {stem} | icons/{stem}.png | {svg_cell} |")
    return "# Icon Library\n\n" + (
        "Each icon is a PNG + SVG pair. When swapping an icon in a slide, replace both files. "
        "See `workflows.md` for the swap procedure.\n\n"
        f"**Total: {len(pngs)} icons.**\n\n"
        "| Name | PNG | SVG |\n|---|---|---|\n" + "\n".join(rows) + "\n"
    )


def render_index_md(per_category_counts: dict[str, int], total: int) -> str:
    lines = [
        "# Slalom PPTX Reference Index",
        "",
        f"Catalog of {total} slide layouts and 80 icons. Read only the section you need; "
        "each per-category file lists ~5-15 slides.",
        "",
        "## Quick lookup",
        "",
        "| Need | Read |",
        "|---|---|",
        "| Brand colors, fonts, footer | `references/colors.md` |",
        "| Icon catalog (PNG/SVG paths) | `references/icons.md` |",
        "| How to assemble a deck (Workflows A/B/C, icon swap, validation) | `references/workflows.md` |",
        "| Choosing a layout (cross-slide decision aid) | `references/selection.md` |",
        "| Programmatic catalog (JSON) | `references/index.json` |",
        "",
        "## Slide layouts by category",
        "",
        "| Category | Read | Slides |",
        "|---|---|--:|",
    ]
    for cat in CATEGORIES:
        n = per_category_counts.get(cat, 0)
        if n == 0:
            continue
        lines.append(f"| {CATEGORY_LABELS[cat]} | `references/slides/{cat}.md` | {n} |")
    lines.extend([
        "",
        "## Asset paths (relative to skill root)",
        "",
        "- Master template (single source of truth): `assets/master/Slalom Template document.pptx`",
        "- Slide extraction: `rearrange.py master.pptx out.pptx <idx>` where `idx` is **0-based** (slide N in PowerPoint = `idx` N-1)",
        "- Text inventories: `assets/inventories/<name>.json`",
        "- Icons: `assets/icons/<name>.png` and `assets/icons/<name>.svg`",
        "- Thumbnails: rendered on demand via LibreOffice (see `references/workflows.md` *On-demand thumbnails*); not pre-baked.",
        "- Reprocess script: `scripts/preprocess.sh`",
        "",
    ])
    return "\n".join(lines)


def render_category_md(
    category: str,
    entries: list[SlideEntry],
    inv_dir: Path,
) -> str:
    label = CATEGORY_LABELS[category]
    lines = [
        f"# {label}",
        "",
        f"{len(entries)} layouts in this category. Pick by matching the message and content shape. "
        "`idx` is the **0-based** master-slide position — pass it directly to "
        "`rearrange.py`. (Slide N as numbered in PowerPoint = `idx` N-1.) "
        "Render any layout's thumbnail on demand via the snippet in `references/workflows.md`.",
        "",
        "| Name | When to use | idx | Key shapes |",
        "|---|---|--:|---|",
    ]
    for e in entries:
        inv = load_inventory(inv_dir / f"{e.name}.json")
        shapes = ", ".join(summarise_shapes(inv)) if inv else "(no inventory yet)"
        lines.append(f"| `{e.name}` | {e.when_to_use} | {e.idx} | {shapes} |")
    lines.append("")
    return "\n".join(lines)


def render_index_json(
    master_pptx: Path,
    inv_dir: Path,
    icons_dir: Path,
    brand_colors: list[tuple[str, str, str, str]],
    vendored_legacy_commit: str | None,
    vendored_modern_commit: str | None,
) -> str:
    slides_payload = []
    for e in SLIDE_MAP:
        inv = load_inventory(inv_dir / f"{e.name}.json")
        slides_payload.append({
            "name": e.name,
            "idx": e.idx,
            "category": e.category,
            "when_to_use": e.when_to_use,
            "key_shapes": summarise_shapes(inv) if inv else [],
            "inventory_path": f"assets/inventories/{e.name}.json",
        })

    icons_payload = []
    for png in sorted(icons_dir.glob("*.png")):
        svg = icons_dir / f"{png.stem}.svg"
        icons_payload.append({
            "name": png.stem,
            "png_path": f"assets/icons/{png.name}",
            "svg_path": f"assets/icons/{svg.name}" if svg.exists() else None,
        })

    payload = {
        "schema": "slalom-pptx-index/v1",
        "master_sha256": compute_master_sha256(master_pptx),
        "master_path": "assets/master/Slalom Template document.pptx",
        "slides_idx_convention": "0-based",
        "vendored_from": {
            "legacy_pptx_skill_commit": vendored_legacy_commit,
            "modern_pptx_skill_commit": vendored_modern_commit,
            "patches_to_legacy": [
                "#5 chmod after copy in rearrange.py (defends against read-only master)",
                "#6 inventory.py overflow detector now prefers layout defRPr over master bodyStyle",
                "#7 replace.py text_frame.clear() gated on shape membership in replacements",
                "#8 replace.py splits run text on \\x0b and emits <a:br/> so soft line breaks survive serialization (no more _x000B_ literals)",
                "#9 inventory.py no longer filters SLIDE_NUMBER placeholders; they are editable per Workflow A's contract",
                "#10 replace.py auto-fills SLIDE_NUMBER placeholders with the slide's actual deck position when the JSON omits them",
                "#11 replace.py overflow errors now name the placeholder_type and include a 30-char preview of the master text",
            ],
            "vendor_path": "scripts/vendor",
        },
        "brand": {
            "colors": [
                {"scheme": s, "hex": f"#{h}", "slalom_name": n, "role": r}
                for s, h, n, r in brand_colors
            ],
            "font_body": "+mn-lt",
            "font_heading": "+mj-lt",
            "footer": "Slalom. All Rights Reserved. Proprietary and Confidential.",
            "date_format": "yyyy",
            "date_max_chars": 4,
            "body_slide_required_fills": {
                "DATE": "deck year in brand.date_format (master DATE placeholder is ~0.28\" wide; longer formats wrap or clip)",
                "FOOTER": "exact text from brand.footer",
                "SLIDE_NUMBER": "slide's actual deck position as plain text (master ships hardcoded; not a dynamic field)",
            },
        },
        "slides": slides_payload,
        "icons": icons_payload,
    }
    return json.dumps(payload, indent=2) + "\n"


def render_all(
    master_pptx: Path,
    inv_dir: Path,
    icons_dir: Path,
    references_dir: Path,
    vendored_legacy_commit: str | None = None,
    vendored_modern_commit: str | None = None,
) -> None:
    references_dir.mkdir(parents=True, exist_ok=True)
    slides_ref_dir = references_dir / "slides"
    slides_ref_dir.mkdir(parents=True, exist_ok=True)

    brand_colors = read_brand_colors(master_pptx)
    grouped = by_category()

    (references_dir / "colors.md").write_text(render_colors_md(brand_colors))
    (references_dir / "workflows.md").write_text(render_workflows_md())
    (references_dir / "icons.md").write_text(render_icons_md(icons_dir))

    for category in CATEGORIES:
        entries = grouped.get(category, [])
        if not entries:
            # Still emit an empty stub so the file always exists; agent can see "(empty)"
            stub = (
                f"# {CATEGORY_LABELS[category]}\n\n"
                "_No layouts extracted in this category yet._\n"
            )
            (slides_ref_dir / f"{category}.md").write_text(stub)
            continue
        (slides_ref_dir / f"{category}.md").write_text(
            render_category_md(category, entries, inv_dir)
        )

    counts = {c: len(grouped.get(c, [])) for c in CATEGORIES}
    total = sum(counts.values())
    (references_dir / "index.md").write_text(render_index_md(counts, total))
    (references_dir / "index.json").write_text(
        render_index_json(
            master_pptx, inv_dir, icons_dir, brand_colors,
            vendored_legacy_commit, vendored_modern_commit,
        )
    )


if __name__ == "__main__":
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument("--master", type=Path, required=True)
    p.add_argument("--inventories", type=Path, required=True)
    p.add_argument("--icons", type=Path, required=True)
    p.add_argument("--references", type=Path, required=True)
    p.add_argument(
        "--vendored-legacy-commit",
        default=None,
        help="Commit hash that the vendored legacy/ scripts (rearrange/replace/inventory) were copied from.",
    )
    p.add_argument(
        "--vendored-modern-commit",
        default=None,
        help="Commit hash that the vendored modern/ tree (clean.py + office/) was copied from.",
    )
    args = p.parse_args()
    render_all(
        args.master, args.inventories, args.icons, args.references,
        vendored_legacy_commit=args.vendored_legacy_commit,
        vendored_modern_commit=args.vendored_modern_commit,
    )
    print(f"Rendered references to {args.references}")
